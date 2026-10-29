# -*- coding: utf-8 -*-
"""İdare şablonuna aktarma: ikili çizelge şablonu seçilen ayın listesiyle dolar."""
import datetime
from io import BytesIO

import openpyxl

from nobet_disa_aktar import fill_idare_sablonu


def _rows(gun, roller):
    return [{"Tarih": f"{g} Tem", roller[0]: f"A{g}", roller[1]: f"B{g}"} for g in range(1, gun + 1)]


def _ac(veri):
    return openpyxl.load_workbook(BytesIO(veri)).active


def test_sablon_ayin_gunleri_ve_adlariyla_dolar():
    roller = ["Yoğun Bakım", "Servis"]
    ws = _ac(fill_idare_sablonu(_rows(31, roller), roller, 2026, 7))
    assert ws["A3"].value == "AY/YIL: TEMMUZ/2026"
    assert ws["A5"].value == datetime.datetime(2026, 7, 1)
    assert ws["B5"].value == "ÇARŞAMBA"
    assert (ws["D5"].value, ws["E5"].value) == ("A1", "B1")
    assert ws["A35"].value == datetime.datetime(2026, 7, 31)
    assert (ws["D35"].value, ws["E35"].value) == ("A31", "B31")
    assert ws["B36"].value == "BİRİM SORUMLUSU"


def test_kisa_ayda_fazla_satirlar_bos_kalir_ve_formul_calismaz():
    roller = ["Yoğun Bakım", "Servis"]
    rows = _rows(28, roller)
    rows[0][roller[0]] = "=1+1"
    rows[1][roller[1]] = "-"
    ws = _ac(fill_idare_sablonu(rows, roller, 2026, 2))
    assert ws["A3"].value == "AY/YIL: ŞUBAT/2026"
    assert ws["D5"].value == "=1+1" and ws["D5"].data_type == "s"
    assert ws["E6"].value is None
    assert ws["A33"].value is None and ws["D34"].value is None
