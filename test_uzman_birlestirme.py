# -*- coding: utf-8 -*-
"""Uzman Excel'i ile ortak liste ve uzman idare şablonu."""
from io import BytesIO

import openpyxl
import pandas as pd

from nobet_disa_aktar import fill_idare_sablonu, ortak_liste_excel, read_uzman_excel
from nobet_gunler import build_gun_detaylari


def _uzman_excel(tarihler, baslik=("Tarih", "Nöbetçi 1", "Nöbetçi 2")):
    cikti = BytesIO()
    pd.DataFrame(
        [(t, f"A{i}", f"B{i}") for i, t in enumerate(tarihler, 1)], columns=list(baslik)
    ).to_excel(cikti, index=False)
    cikti.seek(0)
    return cikti


def test_uzman_excel_secili_ayin_gunlerini_okur_baska_ayi_atlar():
    dosya = _uzman_excel(["01.08.2026", "02.08.2026", "01.09.2026", "kötü"])
    gunler, atlanan = read_uzman_excel(dosya, 2026, 8)
    assert gunler == {1: ("A1", "B1"), 2: ("A2", "B2")}
    assert atlanan == 2


def test_uzman_excel_tarih_sutunu_yoksa_none():
    assert read_uzman_excel(_uzman_excel(["01.08.2026"], ("Gün", "N1", "N2")), 2026, 8) is None


def test_ortak_liste_uzman_ve_kendi_adlari_yan_yana_hafta_sonu_gri():
    detay = build_gun_detaylari(2026, 8, 31, [])
    rows = [{"AYB": f"K{g}", "GYB": f"L{g}"} for g in range(1, 32)]
    ws = openpyxl.load_workbook(BytesIO(
        ortak_liste_excel(rows, ["AYB", "GYB"], {1: ("=A", "B")}, 2026, 8, detay)
    )).active
    assert [c.value for c in ws[1][:5]] == ["TARİH", "UZMAN", None, "AYB", "GYB"]
    assert [c.value for c in ws[2][:5]] == ["01.08.2026 Cmt", "=A", "B", "K1", "L1"]
    assert ws["B2"].data_type == "s"
    assert ws["A2"].fill.fgColor.rgb.endswith("BFBFBF")      # 1 Ağustos Cumartesi
    assert ws["A4"].fill.fgColor.rgb != ws["A2"].fill.fgColor.rgb  # 3 Ağustos Pazartesi
    assert ws["B3"].value is None                             # uzman girilmemiş gün


def test_uzman_idare_sablonu_basligi_degisir():
    rows = [{"Uzman 1": "A", "Uzman 2": "B"}]
    ws = openpyxl.load_workbook(BytesIO(
        fill_idare_sablonu(rows, ["Uzman 1", "Uzman 2"], 2026, 8, baslik="UZMAN BAŞLIK")
    )).active
    assert ws["B1"].value == "UZMAN BAŞLIK"
    assert (ws["D5"].value, ws["E5"].value) == ("A", "B")


from nobet_core import apply_app_settings, collect_app_settings
from nobet_disa_aktar import kisa_adlar, personel_oku, personel_yaz

_PERSONEL = [
    ("Murat", "Özgümüş", "Uzman"), ("Murat", "Alpaslan", "Uzman"), ("Doğuş", "Uyanık", "Uzman"),
    ("Serpil Canan", "Erbuyun", "Uzman"), ("Sinem", "Aksoy", "Uzman"), ("Ali", "Alagöz", "Pratisyen"),
]


def test_kisa_adlar_ilk_ad_ve_ayni_adda_soyad_bas_harfi():
    sonuc = kisa_adlar(_PERSONEL, ["MURAT OZGUMUS", "MURAT ALPASLAN", "DOGUS UYANIK",
                                   "SERPIL CANAN ERBUYUN", "SINEM A", "ALI ALAGOZ", "BILINMEYEN KISI"])
    assert sonuc == {
        "MURAT OZGUMUS": "MURAT Ö.", "MURAT ALPASLAN": "MURAT A.", "DOGUS UYANIK": "DOĞUŞ",
        "SERPIL CANAN ERBUYUN": "SERPİL", "SINEM A": "SİNEM", "ALI ALAGOZ": "ALİ",
        "BILINMEYEN KISI": "BILINMEYEN KISI",
    }


def test_belirsiz_ad_eslesmez():
    kisiler = [("Murat", "Özgümüş", ""), ("Murat", "Özgür", "")]
    assert kisa_adlar(kisiler, ["MURAT O"]) == {"MURAT O": "MURAT O"}


def test_personel_metni_gidip_gelir_ve_noktalı_virgul_temizlenir():
    metin = personel_yaz(_PERSONEL + [("Boş", "", ""), ("A;B", "C", "Uzman")])
    assert personel_oku(metin) == [_p + ("",) for _p in _PERSONEL] + [("A B", "C", "Uzman", "")]


def test_personel_listesi_ayarlarla_kaydedilir():
    durum = {"personel_text": "Ali;Veli;Uzman"}
    ayarlar = collect_app_settings(durum)
    assert ayarlar["personnel"] == "Ali;Veli;Uzman"
    yeni = {}
    apply_app_settings(yeni, ayarlar)
    assert yeni["personel_text"] == "Ali;Veli;Uzman"


def test_kisa_adlar_yazim_farki_ve_eksik_sozcuk():
    kisiler = [("İsmahangül", "Özbayrak", ""), ("Murat", "Alparslan", ""), ("Murat", "Özgümüş", ""),
               ("Arda", "Balcan", ""), ("İrem", "Gür Türkan", ""), ("Sinem", "Akkaya", "")]
    sonuc = kisa_adlar(kisiler, ["ISMIHAN OZBAYRAK", "MURAT ALPASLAN", "MEHMET ARDA BALCAN", "IREM GUR"])
    assert sonuc == {"ISMIHAN OZBAYRAK": "İSMAHANGÜL", "MURAT ALPASLAN": "MURAT A.",
                     "MEHMET ARDA BALCAN": "ARDA", "IREM GUR": "İREM"}


def test_elle_yazilan_kisa_ad_kullanilir_ve_kaydedilir():
    kisiler = [("Murat", "Özgümüş", "Uzman", "MURO"), ("Murat", "Alpaslan", "Uzman", "")]
    assert kisa_adlar(kisiler, ["MURAT OZGUMUS", "MURAT ALPASLAN"]) == {
        "MURAT OZGUMUS": "MURO", "MURAT ALPASLAN": "MURAT A."}
    metin = personel_yaz(kisiler)
    assert metin == "Murat;Özgümüş;Uzman;MURO\nMurat;Alpaslan;Uzman"
    assert personel_oku(metin) == kisiler


def test_ayni_kisa_ada_denk_gelen_iki_kisi_birlesmez():
    kisiler = [("Murat", "Özgümüş", ""), ("Murat", "Özkan", ""), ("Ali", "Yılmaz", "")]
    sonuc = kisa_adlar(kisiler, ["MURAT OZGUMUS", "MURAT OZKAN"])
    assert sonuc["MURAT OZGUMUS"] != sonuc["MURAT OZKAN"]
    # Listede olmayan "ALI YILMAZER" ile "ALI YILMAZ" aynı kişiye bağlanmaz.
    sonuc = kisa_adlar(kisiler, ["ALI YILMAZ", "ALI YILMAZER"])
    assert sonuc == {"ALI YILMAZ": "ALI YILMAZ", "ALI YILMAZER": "ALI YILMAZER"}

def test_uzman_excel_buyuk_harfli_tarih_basligi_okunur():
    gunler, _ = read_uzman_excel(_uzman_excel(["01.08.2026"], ("TARİH", "N1", "N2")), 2026, 8)
    assert gunler == {1: ("A1", "B1")}


def test_uzman_excel_ustte_baslik_satiri_olan_dosya_okunur():
    cikti = BytesIO()
    pd.DataFrame(
        [["UZMAN NÖBET ÇİZELGESİ", None, None], ["TARİH", "AD", "AD"], ["01.08.2026", "A1", "B1"]]
    ).to_excel(cikti, index=False, header=False)
    cikti.seek(0)
    gunler, _ = read_uzman_excel(cikti, 2026, 8)
    assert gunler == {1: ("A1", "B1")}


def test_uzman_excel_iki_ay_alt_alta_ise_ikinci_ay_da_okunur():
    tarihler = [f"{g:02d}.11.2026" for g in range(1, 31)] + [f"{g:02d}.12.2026" for g in range(1, 32)]
    gunler, _ = read_uzman_excel(_uzman_excel(tarihler), 2026, 12)
    assert len(gunler) == 31


def test_idare_sablonu_gun_adlari_turkce_harfli():
    ws = openpyxl.load_workbook(BytesIO(fill_idare_sablonu([], ["A", "B"], 2026, 8))).active
    assert (ws["B5"].value, ws["B6"].value) == ("CUMARTESİ", "PAZAR")  # 1 Ağustos 2026 Cumartesi
    assert ws["B7"].value == "PAZARTESİ"


def _kayitli_liste_uygulamasi():
    import nobet_disa_aktar as m
    from nobet_gunler import build_gun_detaylari as _gd
    gd = _gd(2026, 11, 30, [])
    m.render_uzman_birlestirme([{"AYB": "Ali", "GYB": "Can"}] * 30, ["AYB", "GYB"], 2026, 11, gd,
                               owner="ben@example.com", db_ready=True)


def test_kayitli_uzman_listesi_ay_acilinca_dosyasiz_hazir_gelir(monkeypatch):
    import nobet_disa_aktar
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(nobet_disa_aktar, "load_uzman_list",
                        lambda owner, yil, ay: ("uzman_kasim.xlsx", {1: ("A1", "B1"), 2: ("A2", "")}))
    at = AppTest.from_function(_kayitli_liste_uygulamasi, default_timeout=60)
    at.run()
    assert [e.value for e in at.exception] == []
    assert any("Uzman listesi: uzman_kasim.xlsx yüklendi" in c.value for c in at.caption)
    etiketler = [b.label for b in at.get("download_button")]
    assert "📋 Ortak liste (Excel)" in etiketler and "🏥 Uzman idare şablonu" in etiketler


def _kaydetme_uygulamasi():
    import streamlit as st
    import nobet_disa_aktar as m
    st.session_state.setdefault("kayitlar", [])

    class Dosya:
        file_id, name = "f1", "uzman.xlsx"
    m._uzman_listesini_kaydet("ben", 2026, 11, Dosya, {1: ("A", "B")})
    st.write(len(st.session_state.kayitlar))


def _sahte_kayit(*a):
    import streamlit as st
    st.session_state.kayitlar.append(a)
    return True


def test_yuklenen_uzman_listesi_aya_bir_kez_kaydedilir(monkeypatch):
    import nobet_disa_aktar
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(nobet_disa_aktar, "save_uzman_list", _sahte_kayit)
    at = AppTest.from_function(_kaydetme_uygulamasi, default_timeout=60)
    at.run()
    at.run()  # yeniden çizim tekrar yazmaz
    assert at.session_state["kayitlar"] == [("ben", 2026, 11, "uzman.xlsx", {1: ("A", "B")})]
    assert at.session_state["uzman_listesi_kayit"][("ben", 2026, 11)] == ("uzman.xlsx", {1: ("A", "B")})


def test_kilitli_ayda_uzman_listesi_kaydedilmez(monkeypatch):
    import nobet_disa_aktar
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(nobet_disa_aktar, "save_uzman_list", _sahte_kayit)
    at = AppTest.from_function(_kaydetme_uygulamasi, default_timeout=60)
    at.session_state["ay_kilitli"] = True
    at.run()
    assert at.session_state["kayitlar"] == []
    assert any("Ay kilitli" in i.value for i in at.info)
