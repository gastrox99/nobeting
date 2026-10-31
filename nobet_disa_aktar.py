# -*- coding: utf-8 -*-
"""Listenin altındaki düğme satırı (Dışa aktar, Paylaş, Yedekle) ve Excel yükleme.

PNG/Excel/HTML dosyaları yalnızca düğmeye tıklanınca üretilir.
"""
import calendar
import datetime
import difflib
import importlib.util
import re
import zipfile
from pathlib import Path
from urllib.parse import quote
from io import BytesIO

import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st

from db import create_share_link, delete_all_share_links, get_share_link, log_unexpected
from nobet_onkontrol import render_share_check
from nobet_core import (
    EXCEL_IMPORT_MAX_BYTES,
    EXCEL_IMPORT_MAX_ROWS,
    TR_AY_ADLARI,
    build_preference_share_text,
    build_share_text,
    csv_safe_frame,
    import_schedule_table,
    parse_schedule_import_date,
    preference_overview,
)

# Excel export
EXCEL_AVAILABLE = importlib.util.find_spec("xlsxwriter") is not None


def convert_df_to_png(df):
    fig, ax = plt.subplots(figsize=(8, len(df) * 0.4 + 1))
    try:
        ax.axis('off')
        table = ax.table(cellText=df.values, colLabels=df.columns, loc='center', cellLoc='center')
        table.auto_set_font_size(False); table.set_fontsize(10); table.scale(1.2, 1.2)

        if "Tarih" in df.columns:
            idx = df.columns.get_loc("Tarih")
            for i, row_data in enumerate(df['Tarih']):
                color = "white"
                if "Cmt" in row_data or "Paz" in row_data: color = "#dbeafe"
                for j in range(len(df.columns)):
                    table[i+1, j].set_facecolor(color)
                    if j == idx: table[i+1, j].set_text_props(ha='left')

        for j in range(len(df.columns)):
            table[0, j].set_facecolor("#dddddd")
            table[0, j].set_text_props(weight='bold')

        buf = BytesIO()
        fig.savefig(buf, format='png', bbox_inches='tight', dpi=150)
        return buf.getvalue()
    finally:
        plt.close(fig)

def convert_df_to_excel(df_liste, df_stats_load, df_stats_finance, yil=None):
    """Convert dataframes to Excel with multiple sheets.

    Include the year in exported schedule labels when it is known so a later
    import can reject a schedule from the same month in the wrong year.
    """
    output = BytesIO()
    _excel_list = df_liste.copy()
    if yil is not None and "Tarih" in _excel_list.columns:
        def _add_year_to_date_label(value):
            _parsed = parse_schedule_import_date(value)
            if not _parsed or _parsed["month"] is None or _parsed["year"] is not None:
                return value
            _match = re.match(r"^(\d{1,2}\s+[^\W\d_]+)(.*)$", str(value).strip())
            return f"{_match.group(1)} {yil}{_match.group(2)}" if _match else value

        _excel_list["Tarih"] = _excel_list["Tarih"].map(_add_year_to_date_label)

    # "=" ile başlayan isim/görev adı Excel'de formül olarak çalışmasın.
    with pd.ExcelWriter(output, engine='xlsxwriter',
                        engine_kwargs={"options": {"strings_to_formulas": False}}) as writer:
        _excel_list.to_excel(writer, sheet_name='Günlük Liste', index=False)
        df_stats_load.to_excel(writer, sheet_name='Nöbet Yükü')
        df_stats_finance.to_excel(writer, sheet_name='Ücret Özeti')
        
        def _xl_col(n):
            s = ""
            while n >= 0:
                s = chr(65 + n % 26) + s
                n = n // 26 - 1
            return s
        n_cols = max(
            df_liste.shape[1],
            df_stats_load.shape[1] + 1,
            df_stats_finance.shape[1] + 1,
        )
        last_col = _xl_col(n_cols + 2)
        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]
            worksheet.set_column(f'A:{last_col}', 15)
    
    return output.getvalue()


# Uzman dosyasında birkaç ay alt alta olabilir; bir ay en fazla 31 satır olduğundan 400 satır yeter.
UZMAN_EXCEL_MAX_ROWS = 400
IDARE_SABLONU = Path(__file__).with_name("sablon") / "idare_ikili.xlsx"
_IDARE_GUNLERI = ("PAZARTESİ", "SALI", "ÇARŞAMBA", "PERŞEMBE", "CUMA", "CUMARTESİ", "PAZAR")


def fill_idare_sablonu(liste_rows, role_names, yil, ay, baslik=None):
    """İdarenin ikili nöbet çizelgesi şablonunu seçilen ayın listesiyle doldurur.

    Şablonda ad sütunları D ve E'dir; ilk iki görev bunlara yazılır. Ay 31 günden
    kısaysa fazla satırlar boş kalır. ``baslik`` verilirse şablonun başlığı değişir.
    """
    wb = openpyxl.load_workbook(IDARE_SABLONU)
    ws = wb.active
    if baslik:
        ws["B1"] = baslik
    _ay_adi = TR_AY_ADLARI[ay].replace("i", "İ").replace("ı", "I").upper()
    ws["A3"] = f"AY/YIL: {_ay_adi}/{yil}"
    for gun in range(1, calendar.monthrange(yil, ay)[1] + 1):
        _tarih = datetime.date(yil, ay, gun)
        _satir = gun + 4
        ws[f"A{_satir}"] = _tarih
        ws[f"B{_satir}"] = _IDARE_GUNLERI[_tarih.weekday()]
        _kayit = liste_rows[gun - 1] if gun <= len(liste_rows) else {}
        for _sutun, _gorev in zip("DE", role_names[:2]):
            _ad = _kayit.get(_gorev)
            _hucre = ws[f"{_sutun}{_satir}"]
            if _ad not in (None, "", "-"):
                _hucre.value = str(_ad)
                # "=" ile başlayan ad Excel'de formül olarak çalışmasın.
                _hucre.data_type = "s"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def read_uzman_excel(dosya, yil, ay):
    """Diğer ekibin Excel'i ("Tarih" + iki nöbetçi sütunu): ({gün: (ad1, ad2)}, atlanan satır).

    Seçili ay dışındaki ya da okunamayan tarihli satırlar atlanır (sayılır). "Tarih"
    sütunu ve ardından iki sütun yoksa None döner.
    """
    tablo = pd.read_excel(dosya, header=None, nrows=UZMAN_EXCEL_MAX_ROWS)
    # Başlık ("Tarih", büyük/küçük harf fark etmez) ilk 10 satırda aranır.
    baslik_i = tarih_i = None
    for _i in range(min(10, len(tablo))):
        _sutun = next((_j for _j, _c in enumerate(tablo.iloc[_i]) if "TARIH" in _ascii_buyuk(_c)), None)
        if _sutun is not None:
            baslik_i, tarih_i = _i, _sutun
            break
    if tarih_i is None or tarih_i + 2 >= tablo.shape[1]:
        return None
    gunler, atlanan = {}, 0
    for _, satir in tablo.iloc[baslik_i + 1:].iterrows():
        _t = parse_schedule_import_date(satir.iloc[tarih_i])
        if not _t or (_t["month"], _t["year"]) != (ay, yil) or _t["day"] in gunler:
            atlanan += 1
            continue
        gunler[_t["day"]] = tuple(
            "" if pd.isna(_v) else str(_v).strip() for _v in satir.iloc[tarih_i + 1:tarih_i + 3]
        )
    return gunler, atlanan


_ASCII_HARF = str.maketrans("İıŞşĞğÜüÖöÇç", "IISSGGUUOOCC")


def _ascii_buyuk(metin):
    """Türkçe harfleri ASCII'ye çevirip büyütür ("Murat Özgümüş" → "MURAT OZGUMUS")."""
    return " ".join(str(metin).translate(_ASCII_HARF).upper().split())


def _buyuk(metin):
    return metin.replace("i", "İ").replace("ı", "I").upper()


def personel_oku(metin):
    """"Ad;Soyad;Grup;Kısa ad" satırlarından [(ad, soyad, grup, kısa ad)]; adı ya da soyadı boş satırlar atlanır."""
    kisiler = []
    for satir in (metin or "").splitlines():
        ad, soyad, grup, kisa = ([_p.strip() for _p in satir.split(";")] + ["", "", "", ""])[:4]
        if ad and soyad:
            kisiler.append((ad, soyad, grup, kisa))
    return kisiler


def personel_yaz(kisiler):
    """[(ad, soyad, grup, kısa ad)] → kayıt metni; ad ve soyadı boş olanlar yazılmaz."""
    temiz = [[str(_p or "").replace(";", " ").strip() for _p in k] for k in kisiler]
    return "\n".join(";".join(k).rstrip(";") for k in temiz if k[0] and k[1])


def _eslesen_kisi(kayitlar, dosya_adi):
    """Dosyadaki ada karşılık gelen tek kayıt; birden çok ya da hiç aday varsa None."""
    _ad = _ascii_buyuk(dosya_adi)
    _parca = _ad.split()
    _asamalar = (
        lambda k: k["tam"] == _ad,
        lambda k: k["tam"].startswith(_ad + " "),
        lambda k: len(_parca) > 1 and len(_parca[-1]) == 1
        and k["ilk"] == _parca[0] and k["soyad"][0] == _parca[-1],
        lambda k: _ad.endswith(" " + k["soyad"]),
        lambda k: difflib.SequenceMatcher(None, _ad, k["tam"]).ratio() >= 0.85,
    )
    for _asama in _asamalar:
        _adaylar = [_k for _k in kayitlar if _asama(_k)]
        if _adaylar:
            return _adaylar[0] if len(_adaylar) == 1 else None
    return None


def _kisa_kayitlari(kisiler):
    _ilk = [_ascii_buyuk(_k[0]).split()[0] for _k in kisiler]
    _sayac = {_i: _ilk.count(_i) for _i in _ilk}
    _kayitlar = []
    for _sira, (_k, _i) in enumerate(zip(kisiler, _ilk)):
        ad, soyad = _k[0], _k[1]
        _elle = _k[3].strip() if len(_k) > 3 else ""
        _kisa = _elle or _buyuk(ad.split()[0]) + (f" {_buyuk(soyad)[0]}." if _sayac[_i] > 1 else "")
        _kayitlar.append({"sira": _sira, "ilk": _i, "soyad": _ascii_buyuk(soyad),
                          "tam": _ascii_buyuk(f"{ad} {soyad}"), "kisa": _kisa})
    return _kayitlar


def kisa_adlar(kisiler, dosya_adlari):
    """Dosyadaki adlar ("MURAT OZGUMUS", "SINEM A") → ortak listedeki kısa ad.

    Kısa ad, personel listesinde elle yazılmışsa o; yoksa adın ilk sözcüğüdür; aynı ilk ada sahip başka biri varsa soyadın baş harfi
    eklenir ("MURAT Ö."). Eşleştirme sırayla: tam ad, adın başı ("IREM GUR"), ad + soyad
    baş harfi ("SINEM A"), soyad sonu ("MEHMET ARDA BALCAN"), yazım benzerliği (≥ 0,85).
    Bir aşamada birden çok kişi çıkarsa, hiçbiri çıkmazsa ya da iki dosya adı aynı kısa ada denk gelirse ad olduğu gibi kalır.
    """
    _kayitlar = _kisa_kayitlari(kisiler)
    sonuc = {}
    for _dosya_adi in dosya_adlari:
        _k = _eslesen_kisi(_kayitlar, _dosya_adi)
        sonuc[_dosya_adi] = _k["kisa"] if _k else _dosya_adi
    # İki dosya adı aynı kısa ada çıkıyorsa ikisi de bağlanmaz (iki kişi tek ada karışmasın).
    _cakisan = {_k for _k in sonuc.values() if list(sonuc.values()).count(_k) > 1}
    return {_a: (_a if _k in _cakisan and _k != _a else _k) for _a, _k in sonuc.items()}


def ortak_liste_excel(liste_rows, liste_roles, uzman_adlari, yil, ay, gun_detaylari):
    """Tarih | UZMAN (2 ad) | görev 1 | görev 2 biçiminde ortak liste; hafta sonu ve tatil gri."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ortak Liste"
    _kenar = Side(style="thin")
    _cerceve = Border(left=_kenar, right=_kenar, top=_kenar, bottom=_kenar)
    _gri = PatternFill("solid", fgColor="BFBFBF")
    _roller = list(liste_roles[:2])
    ws["A1"], ws["B1"] = "TARİH", "UZMAN"
    ws.merge_cells("B1:C1")
    for _sutun, _rol in zip("DE", _roller):
        ws[f"{_sutun}1"] = str(_rol)
    for gun, _detay in sorted((d["day_num"], d) for d in gun_detaylari.values()):
        _satir = gun + 1
        _kayit = liste_rows[gun - 1] if gun <= len(liste_rows) else {}
        _adlar = list(uzman_adlari.get(gun, ("", ""))) + [_kayit.get(_rol) for _rol in _roller]
        ws[f"A{_satir}"] = _detay["full_date"]
        for _sutun, _ad in zip("BCDE", _adlar):
            if _ad not in (None, "", "-"):
                ws[f"{_sutun}{_satir}"].value = str(_ad)
                # "=" ile başlayan ad Excel'de formül olarak çalışmasın.
                ws[f"{_sutun}{_satir}"].data_type = "s"
        for _sutun in "ABCDE":
            _hucre = ws[f"{_sutun}{_satir}"]
            _hucre.border = _cerceve
            if _detay["weekend"] or _detay["holiday"]:
                _hucre.fill = _gri
    for _sutun in "ABCDE":
        ws[f"{_sutun}1"].font = Font(bold=True)
        ws[f"{_sutun}1"].alignment = Alignment(horizontal="center")
        ws[f"{_sutun}1"].border = _cerceve
        ws.column_dimensions[_sutun].width = 20
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _render_personel_listesi():
    """Uzman ve pratisyenlerin ad-soyad listesi; kısa adlar buradan bulunur, ayarlarla birlikte kaydedilir."""
    with st.expander("📇 Personel listesi"):
        st.caption("Ad, soyad ve grup (Uzman / Pratisyen). Ortak listedeki kısa adlar buradan bulunur; "
                   "Kısa ad boşsa adın ilk sözcüğü kullanılır. Satır ekleyip silerek düzenleyin.")
        _kayit = st.session_state.get("personel_text", "")
        _tablo = pd.DataFrame(personel_oku(_kayit), columns=["Ad", "Soyad", "Grup", "Kısa ad"])
        _yeni = st.data_editor(
            _tablo, num_rows="dynamic", hide_index=True, width="stretch", key="personel_tablosu",
            column_config={"Grup": st.column_config.SelectboxColumn("Grup", options=["Uzman", "Pratisyen"])},
        )
        with st.form("personel_toplu_ekle", clear_on_submit=True):
            _toplu = st.text_area("Toplu ekle (her satıra bir ad soyad)", height=110)
            _grup = st.selectbox("Grup", ["Uzman", "Pratisyen"])
            if st.form_submit_button("Ekle"):
                _var = {_ascii_buyuk(f"{_a} {_s}") for _a, _s, *_ in personel_oku(_kayit)}
                _eklenen = [(" ".join(_p[:-1]), _p[-1], _grup, "") for _p in (_l.split() for _l in _toplu.splitlines())
                            if len(_p) > 1 and _ascii_buyuk(" ".join(_p)) not in _var]
                if _eklenen:
                    st.session_state.personel_text = personel_yaz(personel_oku(_kayit) + _eklenen)
                    st.rerun()
        _yeni_metin = personel_yaz(_yeni.itertuples(index=False, name=None))
        if _yeni_metin != _kayit:
            # Ayarlar bir sonraki çalışmada kaydedilir; hemen yeniden çalıştırıp kaydı beklet.
            st.session_state.personel_text = _yeni_metin
            st.rerun()


def render_uzman_birlestirme(liste_rows, liste_roles, yil, ay, gun_detaylari):
    """Dışa aktar içinde: uzman Excel'i yüklenince ortak liste ve uzman idare şablonu indirilir."""
    _render_personel_listesi()
    with st.expander("👥 Uzman listesi ile birleştir"):
        st.caption("Diğer ekibin Excel'ini yükleyin (Tarih + iki nöbetçi sütunu). Ortak liste ve "
                   "uzman idare şablonu hazırlanır; bu ayın kendi listeniz de aynı yerde yer alır.")
        _dosya = st.file_uploader("Uzman Excel (.xlsx)", type=["xlsx"], key="uzman_excel",
                                  label_visibility="collapsed")
        if _dosya is None:
            return
        if _dosya.size > EXCEL_IMPORT_MAX_BYTES:
            st.error("Dosya çok büyük (en fazla 2 MB).")
            return
        try:
            _okunan = read_uzman_excel(_dosya, yil, ay)
        except (ValueError, OSError, zipfile.BadZipFile):
            st.error("Dosya okunamadı. Excel (.xlsx) olduğundan emin olun.")
            return
        except Exception as _e:
            log_unexpected("Uzman Excel", _e)
            st.error("Dosya okunamadı. Excel (.xlsx) olduğundan emin olun.")
            return
        if _okunan is None:
            st.error("Dosya formatı uyumsuz: 'Tarih' sütunu ve ardından iki nöbetçi sütunu gerekli.")
            return
        _gunler, _atlanan = _okunan
        if not _gunler:
            st.warning(f"Seçili döneme ({ay:02d}/{yil}) ait satır bulunamadı.")
            return
        if len(_gunler) < len(gun_detaylari) or _atlanan:
            st.warning(f"⚠️ {len(gun_detaylari)} günden {len(_gunler)} gün okundu; "
                       f"{_atlanan} satır (başka ay ya da okunamayan tarih) atlandı.")
        _adlar = sorted({_ad for _cift in _gunler.values() for _ad in _cift if _ad})
        _personel = personel_oku(st.session_state.get("personel_text", ""))
        _kisa = kisa_adlar(_personel, _adlar)
        _bulunamayan = [_ad for _ad in _adlar if _kisa[_ad] == _ad]
        if _bulunamayan and st.session_state.get("personel_text"):
            st.warning("⚠️ Personel listesinde bulunamadı ya da başka bir adla karıştı: " + ", ".join(_bulunamayan))
        st.caption("Listede görünecek adları düzeltebilirsiniz (ör. MURAT OZGUMUS → MURAT Ö.).")
        _duzenlenen = st.data_editor(
            pd.DataFrame({"Dosyadaki ad": _adlar, "Listede görünecek ad": [_kisa[_ad] for _ad in _adlar]}),
            disabled=["Dosyadaki ad"], hide_index=True, width="stretch", key=f"uzman_adlar_{_dosya.file_id}",
        )
        _esle = dict(zip(_duzenlenen["Dosyadaki ad"], _duzenlenen["Listede görünecek ad"].fillna("")))
        # Düzeltilen kısa adlar personel listesine kaydedilir; bir dahaki sefere hazır gelir.
        _kayitlar = _kisa_kayitlari(_personel)
        _guncel = [list(_p) for _p in _personel]
        for _ad in _adlar:
            _yeni_kisa = str(_esle[_ad]).strip()
            _k = _eslesen_kisi(_kayitlar, _ad)
            if _k and _yeni_kisa and _yeni_kisa != _kisa[_ad]:
                _guncel[_k["sira"]][3] = _yeni_kisa
        if _guncel != [list(_p) for _p in _personel]:
            st.session_state.personel_text = personel_yaz(_guncel)
            st.rerun()
        _gorunen = {g: tuple(_esle.get(_ad, _ad) for _ad in _cift) for g, _cift in _gunler.items()}
        _baslik = st.text_input("Uzman çizelgesi başlığı", "YOĞUN BAKIM UZMAN DOKTOR NÖBET ÇİZELGESİ")
        _uzman_rows = [
            dict(zip(("Uzman 1", "Uzman 2"), _gorunen.get(g, ("", ""))))
            for g in range(1, len(gun_detaylari) + 1)
        ]
        _xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        st.download_button("📋 Ortak liste (Excel)",
                           ortak_liste_excel(liste_rows, liste_roles, _gorunen, yil, ay, gun_detaylari),
                           f"ortak_liste_{yil}_{ay:02d}.xlsx", _xlsx, on_click="ignore", width="stretch")
        st.download_button("🏥 Uzman idare şablonu",
                           fill_idare_sablonu(_uzman_rows, ["Uzman 1", "Uzman 2"], yil, ay, baslik=_baslik),
                           f"idare_sablonu_uzman_{yil}_{ay:02d}.xlsx", _xlsx, on_click="ignore", width="stretch")


def whatsapp_url(text):
    """WhatsApp'ı verilen metin yazılı açan bağlantı (telefonda uygulama açılır)."""
    return "https://wa.me/?text=" + quote(text, safe="")


def app_base_url():
    try:
        _url = st.context.url
    except Exception:
        _url = None
    if _url:
        return _url.split("?")[0].split("#")[0].rstrip("/")
    try:
        return st.secrets["auth"]["redirect_uri"].rsplit("/oauth2callback", 1)[0]
    except Exception:
        return ""


def render_toolbar(df_liste, df_stats_load, df_stats_finance, liste_rows, liste_roles,
                   yil, ay, gun_detaylari, isimler, sutunlar, role_names,
                   db_ready, current_owner, save_undo_state, share_check=()):
    """Dışa aktar / Paylaş / Yedekle düğmelerini çizer."""
    # PNG/Excel/HTML yalnızca butona tıklanınca üretilir (eskiden her nöbet
    # işaretlemesinde matplotlib ile yeniden çiziliyordu). Varsayılan argümanlar
    # o anki tabloları sabitler; on_click="ignore" indirmede sayfayı yeniden çizmez.
    def _png_bytes(_liste=df_liste.copy()):
        return convert_df_to_png(_liste)

    def _excel_bytes(_liste=df_liste.copy(), _load=df_stats_load.copy(),
                     _finance=df_stats_finance.copy(), _yil=yil):
        return convert_df_to_excel(_liste, _load, _finance, yil=_yil)

    def _idare_bytes(_rows=liste_rows, _roles=list(liste_roles), _yil=yil, _ay=ay):
        return fill_idare_sablonu(_rows, _roles, _yil, _ay)

    _share_cache = st.session_state.setdefault("share_token_cache", {})
    # Excel yükleme seyrek kullanıldığı için Yedek bölümünde (render_excel_import).
    # Kapsayıcı anahtarı (bottom-actions) telefonda dört düğmeyi tek satırda tutar (nobet_stil.py).
    _toolbar = st.container(key="bottom-actions")
    _export_col, _share_col, _pref_col, _snapshot_col, _ = _toolbar.columns([1, 1, 1, 1, 2])
    with _export_col:
        with st.popover("📤 Dışa aktar", width="stretch"):
            st.download_button("📊 Excel", _excel_bytes, f"nobet_{yil}_{ay:02d}.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               on_click="ignore", width="stretch", disabled=not EXCEL_AVAILABLE)
            st.download_button("🏥 İdare şablonuna aktar", _idare_bytes, f"idare_sablonu_{yil}_{ay:02d}.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               on_click="ignore", width="stretch")
            if len(liste_roles) > 2:
                st.caption("⚠️ İdare şablonunda yalnız ilk iki görev sütunu var; 3. ve sonraki "
                           "kişiler şablona yazılmaz.")
            render_uzman_birlestirme(liste_rows, liste_roles, yil, ay, gun_detaylari)
            # Sayfanın kendi yazdırma penceresi (A4 yatay, yalnız tablo ve takvim; nobet_stil.py).
            st.iframe(
                "<button onclick=\"window.parent.print()\" style=\"width:100%;padding:8px;border-radius:8px;"
                "border:1px solid rgba(128,128,128,.4);background:transparent;color:inherit;font-size:14px;"
                "cursor:pointer;\">🖨️ Yazdır (A4)</button>",
                height=48,
            )
            st.download_button("🖼️ Resim (PNG)", _png_bytes, f"nobet_{yil}_{ay:02d}.png", "image/png",
                               on_click="ignore", width="stretch")
            # BOM'lu UTF-8: Excel'de Ş, ğ, ı gibi harfler bozulmadan açılır.
            st.download_button("📄 CSV", csv_safe_frame(df_liste).to_csv(index=False).encode('utf-8-sig'),
                               f"nobet_{yil}_{ay:02d}.csv", "text/csv",
                               on_click="ignore", width="stretch")
            render_share_check(share_check)
            st.markdown("**💬 WhatsApp metni**")
            _wa_text = build_share_text(
                liste_rows, liste_roles, yil, ay,
                special_days=[d["day_num"] for d in gun_detaylari.values() if d["weekend"] or d["holiday"]],
            )
            # Telefonda WhatsApp'ı metin yazılı açar; kişi ya da grup orada seçilir.
            st.link_button("💬 WhatsApp'ta gönder", whatsapp_url(_wa_text), type="primary", width="stretch")
            st.caption("Ya da sağ üstteki simgeyle kopyalayıp yapıştırın. Hafta sonu ve tatiller kalın görünür.")
            st.code(_wa_text, language=None)
    with _share_col:
        with st.popover("🔗 Paylaş", width="stretch"):
            if not db_ready:
                st.caption("Paylaşım için veritabanı bağlantısı gerekli.")
            else:
                # Tek kalıcı bağlantı (yıl 0, ay 0): bütün ayları gösterir.
                if "all" not in _share_cache:
                    _share_cache["all"] = get_share_link(current_owner, 0, 0)
                _token = _share_cache["all"]
                render_share_check(share_check)
                st.caption(
                    "Ekip arkadaşlarınız bu tek bağlantıyla giriş yapmadan bütün ayların listesini görebilir; "
                    "her ay yeni bağlantı göndermeniz gerekmez. Kendi adlarını seçip nöbetlerini takvimlerine ekleyebilirler."
                )
                if _token:
                    _share_url = f"{app_base_url()}?paylas={_token}"
                    st.code(_share_url, language=None)
                    st.link_button(
                        "💬 Bağlantıyı WhatsApp'ta gönder",
                        whatsapp_url(f"Nöbet listesi (kendi adını seçip nöbetlerini görebilirsin): {_share_url}"),
                        width="stretch",
                    )
                    if st.button("🔄 Yeni bağlantı oluştur (eskileri çalışmaz)", width="stretch"):
                        # Eski aylık bağlantılar da kapanır.
                        if delete_all_share_links(current_owner):
                            _share_cache["all"] = create_share_link(current_owner, 0, 0)
                        st.rerun()
                    if st.button("🚫 Paylaşımı kapat", width="stretch"):
                        if delete_all_share_links(current_owner):
                            _share_cache["all"] = None
                        st.rerun()
                elif st.button("🔗 Bağlantı oluştur", type="primary", width="stretch"):
                    _share_cache["all"] = create_share_link(current_owner, 0, 0)
                    if not _share_cache["all"]:
                        st.error("Bağlantı oluşturulamadı.")
                    else:
                        st.rerun()

    with _pref_col:
        with st.popover("📝 Tercihler", width="stretch"):
            # Liste hazırlanmadan önce herkesin girilen günlerini grupla teyit etmek için.
            # Onay bekleyen ekip talepleri dahil değil; yalnızca tabloya işlenmiş tercihler.
            _pref_rows = preference_overview(st.session_state.get("pref_df", pd.DataFrame()), isimler)
            if not _pref_rows:
                st.caption("Bu ay Tercih Belirle'de işaretli tercih, kaçınma, izin ya da müsait değil günü yok.")
            else:
                st.caption(
                    f"{TR_AY_ADLARI[ay]} {yil} için herkesin girilen günleri. Listeyi hazırlamadan önce "
                    "grupla paylaşıp yanlış olan varsa düzeltebilirsiniz. Onay bekleyen ekip talepleri dahil değildir."
                )
                _pref_table = pd.DataFrame(_pref_rows)
                st.dataframe(_pref_table, hide_index=True, width="stretch")
                _pref_text = build_preference_share_text(_pref_rows, yil, ay)
                st.link_button("💬 WhatsApp'ta gönder", whatsapp_url(_pref_text),
                               type="primary", width="stretch")

                def _pref_png(_table=_pref_table.copy()):
                    return convert_df_to_png(_table)

                st.download_button("🖼️ Tablo resmi (PNG)", _pref_png, f"tercihler_{yil}_{ay:02d}.png",
                                   "image/png", on_click="ignore", width="stretch")
                st.caption("Ya da sağ üstteki simgeyle kopyalayıp yapıştırın.")
                st.code(_pref_text, language=None)

    with _snapshot_col:
        if db_ready:
            # Tıklama sayfa yeniden çizilmeden önce işlenir; yedek, bu çalışmada
            # son değişiklik kaydedildikten sonra alınır (_run_requested_snapshot).
            st.button(
                "💾 Yedekle", width="stretch", key="snapshot_take_btn",
                help=f"{TR_AY_ADLARI[ay]} {yil} ayını yedekler. Her ayın tek yedeği olur; tekrar basarsanız eskisinin yerine geçer.",
                on_click=lambda: st.session_state.update(snapshot_requested=True),
            )


def render_excel_import(yil, ay, gun_detaylari, isimler, sutunlar, role_names, save_undo_state):
    """Daha önce indirilen Excel listesini bu aya yükler (Yedek bölümünde)."""
    st.caption("Daha önce indirilen Excel dosyasını veya aynı formatta hazırladığınız listeyi yükleyin.")
    for _level, _msg in st.session_state.pop("excel_import_messages", []):
        getattr(st, _level)(_msg)
    _uploaded = st.file_uploader("Excel dosyası seçin (.xlsx)", type=["xlsx"], key="excel_import_uploader", label_visibility="collapsed",
                                  disabled=bool(st.session_state.get("ay_kilitli")))
    # Dosya kutuda kaldıkça sayfa her yenilendiğinde yeniden içe aktarılıp
    # sonsuz döngüye giriyordu (elle yapılan değişiklikler de eziliyordu).
    # Her yüklenen dosya yalnızca bir kez işlenir.
    _upload_id = (
        getattr(_uploaded, "file_id", None) or (_uploaded.name, _uploaded.size)
        if _uploaded is not None else None
    )
    if _uploaded is not None and _upload_id != st.session_state.get("excel_import_done"):
        st.session_state.excel_import_done = _upload_id
        _msgs = []
        _imp = None
        if _uploaded.size > EXCEL_IMPORT_MAX_BYTES:
            # Uygulamanın kendi Excel'i birkaç KB; büyük dosya sunucuyu yorar.
            _msgs.append(("error", "Dosya çok büyük (en fazla 2 MB). Uygulamadan indirilen Excel'i yükleyin."))
        else:
            try:
                _xl = pd.ExcelFile(_uploaded)
                _sheet = "Günlük Liste" if "Günlük Liste" in _xl.sheet_names else _xl.sheet_names[0]
                _imp = import_schedule_table(
                    # Bir ay en fazla 31 gün; fazlası okunmaz.
                    pd.read_excel(_xl, sheet_name=_sheet, nrows=EXCEL_IMPORT_MAX_ROWS),
                    isimler, sutunlar, gun_detaylari, role_names, yil, ay,
                )
            except Exception as _e:
                if not isinstance(_e, (ValueError, OSError, zipfile.BadZipFile)):
                    log_unexpected("Excel içe aktarma", _e)
                _msgs.append(("error",
                    "Dosya okunamadı. Lütfen uygulamadan indirilen Excel formatını "
                    "('Tarih' + görev sütunları) kullandığınızdan emin olun."))
        if _imp is None:
            pass
        elif not _imp["format_ok"]:
            _msgs.append(("error", "Dosya formatı uyumsuz: 'Tarih' sütunu ve en az bir görev sütunu gerekli."))
        elif _imp["date_mismatch"]:
            _msgs.append(("warning",
                f"⚠️ Excel tarihleri seçili dönemle uyuşmuyor ({ay:02d}/{yil}). "
                f"{_imp['date_mismatch']['day']} günü için farklı ay/yıl veya geçersiz gün "
                "bulundu; yükleme iptal edildi."))
        elif _imp["duplicate_day"]:
            _msgs.append(("warning",
                f"⚠️ Excel'de {_imp['duplicate_day']}. gün iki satırda geçiyor; yükleme iptal edildi. "
                "Yinelenen satırı silip tekrar yükleyin."))
        elif _imp["matched"] == 0:
            _msgs.append(("warning",
                "Eşleşen gün bulunamadı; Excel tarihleri okunamadı. Tarih sütununda "
                "gün/ay/yıl bilgilerinin bulunduğunu kontrol edin."
                if _imp["unreadable_rows"] else
                f"Eşleşen gün bulunamadı. Excel tarihleri seçili dönemle ({ay:02d}/{yil}) "
                "uyumlu değil; yükleme iptal edildi."))
        else:
            save_undo_state()
            st.session_state.schedule_bool = _imp["schedule"]
            st.session_state.locked_cells = set()
            st.session_state.cached_rows_liste = _imp["rows"]
            st.session_state.cached_role_names = role_names
            st.session_state.cached_key = (yil, ay, tuple(isimler), tuple(sutunlar), tuple(role_names))
            st.session_state.cached_role_schedule = _imp["schedule"].copy(deep=True)
            st.session_state.should_regenerate_assignments = False
            # Kayıt, sayfa yenilenince otomatik kayıtla yapılır (hata olursa orada görünür).
            _msgs.append(("success", f"{_imp['matched']}/{len(sutunlar)} gün yüklendi."))
            if _imp["skipped_names"]:
                _msgs.append(("warning",
                    f"⚠️ {len(_imp['skipped_names'])} isim mevcut ekipte bulunamadığı için "
                    f"atlandı: {', '.join(_imp['skipped_names'])}"))
            if _imp["matched"] < len(sutunlar):
                _msgs.append(("warning",
                    f"⚠️ Excel'de {len(sutunlar)} günden {_imp['matched']} gün eşleşti; "
                    f"{len(sutunlar) - _imp['matched']} gün boş bırakıldı."))
            if _imp["extra_people"]:
                _msgs.append(("info",
                    f"{_imp['extra_people']} nöbet görev sayısından fazla sütunlardan (Ek 1 gibi) "
                    "okundu; listede 'Ek' sütununda görünür."))
            if _imp["unreadable_rows"]:
                _msgs.append(("warning", f"⚠️ {_imp['unreadable_rows']} satırın tarihi okunamadığı için atlandı."))
        st.session_state.excel_import_messages = _msgs
        st.rerun()
